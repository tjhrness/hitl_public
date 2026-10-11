import io
import zipfile

import pytest
from docx import Document

from inter_rater_analyzer.artifacts import sha256
from inter_rater_analyzer.domain import DomainError
from inter_rater_analyzer.reports import build_model, render_visuals
from tests.integration.test_execution import prepared


def test_immutable_report_package_content_and_privacy(tmp_path):
    services, state = prepared(tmp_path)
    services.execution.execute(state.version_id)
    report = services.reports.generate(state.version_id)
    assert sha256(report.path.read_bytes()) == report.sha256
    with zipfile.ZipFile(report.path) as archive:
        assert {"report.pdf", "report.docx"} <= set(archive.namelist())
        assert any(name.endswith(".png") for name in archive.namelist())
        document = Document(io.BytesIO(archive.read("report.docx")))
        text = "\n".join(p.text for p in document.paragraphs)
        assert "Human review summary" in text
        assert "Audit appendix" in text
        assert "Alice" not in text and "Bob" not in text
        assert archive.read("report.pdf").startswith(b"%PDF")
    with pytest.raises(DomainError):
        services.reports.generate(state.version_id)


def test_report_failure_leaves_no_committed_package(tmp_path, monkeypatch):
    services, state = prepared(tmp_path)
    services.execution.execute(state.version_id)

    def fail(*args, **kwargs):
        raise RuntimeError("private identity")

    monkeypatch.setattr("inter_rater_analyzer.reports.render_documents", fail)
    with pytest.raises(DomainError):
        services.reports.generate(state.version_id)
    assert not list(services.paths.reports.rglob("*.zip"))
    with services.database.connection() as conn:
        assert conn.execute("SELECT count(*) FROM report_artifacts").fetchone()[0] == 0


def test_human_review_is_top_one_and_all_bland_altman_plots(tmp_path):
    services, state = prepared(tmp_path, additional_rows=(("Q2", "Yes", "Yes"),))
    services.execution.execute(state.version_id)
    model = build_model(services.execution, state.version_id)
    human = next(s for s in model.sections if s.title == "Human review summary")
    assert "Q1" in "\n".join(human.paragraphs)
    assert "Q2" not in "\n".join(human.paragraphs)
    snapshot = services.execution.load(state.version_id)
    criterion = snapshot["result"]["criteria"][0]
    criterion["metrics"] = [
        dict(
            estimate=0,
            pair=["Rater A", "Rater B"],
            bland_altman=dict(means=[1, 2], differences=[-1, 1], lower=-2, upper=2),
        )
        for _ in range(3)
    ]
    charts = render_visuals(snapshot, tmp_path)
    assert len(charts) == 4
