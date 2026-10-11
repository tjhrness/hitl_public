from dataclasses import replace

import pytest
from openpyxl import load_workbook

from inter_rater_analyzer.domain import DomainError, RatingType, VersionType
from tests.integration.test_execution import prepared


def test_configuration_reuse_and_numbering(tmp_path):
    services, state = prepared(tmp_path)
    services.execution.execute(state.version_id)
    new = services.versions.create(state.version_id, VersionType.CONFIGURATION_ONLY, "Repeat")
    assert new.workbook_id == state.workbook_id
    assert new.configuration == state.configuration
    assert new.types_locked
    assert services.analyses.get_version(new.version_id).number == 2
    services.drafts.save(replace(new, acknowledged_warnings=state.acknowledged_warnings))
    services.execution.execute(new.version_id)


@pytest.mark.parametrize("change", ["sheet", "item", "rater", "order", "type", "note"])
def test_incompatibility_creates_no_version(tmp_path, change):
    services, state = prepared(tmp_path)
    services.execution.execute(state.version_id)
    source = tmp_path / "source.xlsx"
    book = load_workbook(source)
    sheet = book.active
    if change == "sheet":
        sheet.title = "Different"
    if change == "item":
        sheet["A2"] = "Different"
    if change == "rater":
        sheet["B1"] = "Rater: Carol"
    if change == "order":
        sheet["B1"], sheet["C1"] = "Rater: Bob", "Rater: Alice"
    book.save(source)
    book.close()
    config = state.configuration
    if change == "type":
        config = replace(
            config, criteria=(replace(config.criteria[0], rating_type=RatingType.BINARY),)
        )
    with pytest.raises(DomainError):
        services.versions.create(
            state.version_id,
            VersionType.DATA_CHANGE,
            "" if change == "note" else "Change",
            source=source,
            configuration=config,
        )
    assert services.drafts.active() is None
    with services.database.connection() as conn:
        assert conn.execute("SELECT count(*) FROM analysis_versions").fetchone()[0] == 1


def test_reordered_rows_and_normalized_identities_keep_aliases(tmp_path):
    services, state = prepared(tmp_path, additional_rows=(("Q2", "Yes", "No"),))
    services.execution.execute(state.version_id)
    source = tmp_path / "source.xlsx"
    book = load_workbook(source)
    book.active["B1"] = "Rater: ALICE "
    book.active["B2"] = "No"
    first = [c.value for c in book.active[2]]
    second = [c.value for c in book.active[3]]
    for column, value in enumerate(second, 1):
        book.active.cell(2, column, value)
    for column, value in enumerate(first, 1):
        book.active.cell(3, column, value)
    book.save(source)
    book.close()
    new = services.versions.create(
        state.version_id, VersionType.DATA_CHANGE, "New rating", source=source
    )
    assert new.workbook_id != state.workbook_id
    assert services.drafts.resume_workbook().dataset.criteria[0].raters == ("Rater A", "Rater B")
    assert services.drafts.resume_workbook().dataset.item_ids == ("Q1", "Q2")
    with pytest.raises(DomainError):
        services.drafts.save(replace(new, types_locked=False))
