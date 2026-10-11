from dataclasses import asdict

from openpyxl import load_workbook

from inter_rater_analyzer.services import Services
from tests.workbook_fixtures.builders import workbook


def test_import_aliases_persist_after_restart(tmp_path):
    root = tmp_path / "state"
    services = Services.open(root)
    _, version = services.analyses.create("Study")
    path = workbook(tmp_path / "input.xlsx")
    first = services.workbooks.import_workbook(version.id, path)
    book = load_workbook(path)
    book.active["B1"] = "Rater: ALICE "
    book.save(path)
    book.close()
    services = Services.open(root)
    second = services.workbooks.import_workbook(version.id, path)
    assert first.criteria == second.criteria
    assert "Alice" not in str(asdict(second))
    with services.database.connection() as conn:
        assert (
            conn.execute("SELECT original FROM rater_mappings ORDER BY position").fetchone()[0]
            == "Alice"
        )


def test_invalid_workbook_returns_multiple_findings(tmp_path):
    services = Services.open(tmp_path / "state")
    _, version = services.analyses.create("Study")
    path = workbook(tmp_path / "input.xlsx")
    book = load_workbook(path)
    book.active["B2"] = "=1+1"
    book.active["C2"] = None
    book.save(path)
    book.close()
    result = services.workbooks.import_workbook(version.id, path)
    assert result.dataset is None
    assert {f.code for f in result.findings} >= {"VALIDATION_FORMULA", "VALIDATION_MISSING_RATING"}
