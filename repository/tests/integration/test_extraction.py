import pytest
from openpyxl import load_workbook

from inter_rater_analyzer.domain import DomainError
from inter_rater_analyzer.ingestion import extract_workbook
from tests.workbook_fixtures.builders import workbook


def test_visibility_and_extra_columns(tmp_path):
    path = workbook(tmp_path / "input.xlsx")
    book = load_workbook(path)
    sheet = book.active
    sheet.append(["Q2", "hidden", "hidden"])
    sheet.row_dimensions[3].hidden = True
    sheet["D1"] = "Rater: Secret"
    sheet["D2"] = "=1+1"
    sheet.column_dimensions["D"].hidden = True
    sheet["E1"] = "Comments"
    sheet["E2"] = "=2+2"
    book.save(path)
    book.close()
    criteria = extract_workbook(path)
    assert len(criteria[0].rows) == 1
    assert criteria[0].rows[0].ratings == ("Yes", "No")
    assert criteria[0].rows[0].excel_row == 2


@pytest.mark.parametrize("mode", ["header", "formula", "blank"])
def test_blocking_cells(tmp_path, mode):
    path = workbook(tmp_path / "input.xlsx")
    book = load_workbook(path)
    sheet = book.active
    if mode == "header":
        sheet["A1"] = "item id"
    elif mode == "formula":
        sheet["B2"] = "=1+1"
    else:
        sheet["A2"] = None
    book.save(path)
    book.close()
    with pytest.raises(DomainError):
        extract_workbook(path)


def test_wrong_case_rater_ignored(tmp_path):
    path = workbook(tmp_path / "input.xlsx")
    book = load_workbook(path)
    book.active["C1"] = "rater: Bob"
    book.save(path)
    book.close()
    assert len(extract_workbook(path)[0].rater_headers) == 1
