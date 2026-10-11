"""Workbook inspection without rating normalization or identity disclosure."""

from dataclasses import dataclass, field
from pathlib import Path
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from inter_rater_analyzer.artifacts import sha256
from inter_rater_analyzer.domain import DomainError, ErrorCode, WorkbookArtifactId
from inter_rater_analyzer.findings import ValidationFinding


@dataclass(frozen=True)
class WorksheetMetadata:
    name: str
    rows: int
    columns: int


@dataclass(frozen=True)
class WorkbookInspection:
    sha256: str
    worksheets: tuple[WorksheetMetadata, ...]
    artifact_id: WorkbookArtifactId | None = None


def inspect_workbook(
    path: Path, artifact_id: WorkbookArtifactId | None = None, *, validate_count: bool = True
) -> WorkbookInspection:
    if path.suffix.lower() != ".xlsx":
        raise DomainError(ErrorCode.FILE_INVALID, "Select an .xlsx workbook.")
    try:
        digest = sha256(path.read_bytes())
        book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        try:
            sheets = tuple(
                WorksheetMetadata(sheet.title, sheet.max_row or 0, sheet.max_column or 0)
                for sheet in book.worksheets
                if sheet.sheet_state == "visible"
            )
        finally:
            book.close()
    except (OSError, BadZipFile, InvalidFileException, KeyError, ValueError, ParseError) as error:
        raise DomainError(ErrorCode.FILE_INVALID, "Workbook is unreadable or corrupt.") from error
    if validate_count and not 1 <= len(sheets) <= 10:
        raise DomainError(ErrorCode.VALIDATION_VALUE, "Use 1–10 visible criterion worksheets.")
    return WorkbookInspection(digest, sheets, artifact_id)


@dataclass(frozen=True)
class RawRow:
    excel_row: int
    item_id: object
    ratings: tuple[object, ...]


@dataclass(frozen=True)
class RawCriterion:
    """Internal extraction DTO, never exposed to UI or reports."""

    name: str
    rater_headers: tuple[str, ...] = field(repr=False)
    rows: tuple[RawRow, ...]


def extract_workbook(
    path: Path, findings: list[ValidationFinding] | None = None
) -> tuple[RawCriterion, ...]:
    inspection = inspect_workbook(path, validate_count=findings is None)

    def problem(code: str, message: str, worksheet: str, row: int | None = None) -> None:
        if findings is None:
            raise DomainError(ErrorCode.VALIDATION_VALUE, message)
        findings.append(ValidationFinding(code, message, message, worksheet=worksheet, row=row))

    book = load_workbook(path, data_only=False, keep_links=False)
    criteria: list[RawCriterion] = []
    try:
        visible = [sheet for sheet in book.worksheets if sheet.sheet_state == "visible"]
        for metadata, sheet in zip(inspection.worksheets, visible, strict=True):
            hidden_columns: set[int] = set()
            for dimension in sheet.column_dimensions.values():
                if dimension.hidden:
                    hidden_columns.update(range(dimension.min or 1, (dimension.max or 1) + 1))
            columns = [
                column for column in range(1, sheet.max_column + 1) if column not in hidden_columns
            ]
            last_row = max(sheet.max_row, max(sheet.row_dimensions, default=1))
            rows = [row for row in range(1, last_row + 1) if not sheet.row_dimensions[row].hidden]
            if not columns or not rows or sheet.cell(rows[0], columns[0]).value != "Item ID":
                problem(
                    "VALIDATION_ID_HEADER",
                    "First visible column must have the exact header Item ID.",
                    metadata.name,
                )
                continue
            header_row = rows[0]
            raters: list[tuple[int, str]] = []
            for column in columns[1:]:
                header = sheet.cell(header_row, column).value
                if isinstance(header, str) and header.startswith("Rater:"):
                    raters.append((column, header))
            records: list[RawRow] = []
            for row in rows[1:]:
                item = sheet.cell(row, columns[0])
                if item.value is None or not str(item.value).strip():
                    if findings is None:
                        problem("VALIDATION_ID_BLANK", "Enter an Item ID.", metadata.name, row)
                cells = [sheet.cell(row, column) for column, _ in raters]
                if item.data_type == "f" or any(cell.data_type == "f" for cell in cells):
                    problem(
                        "VALIDATION_FORMULA",
                        f"Use literal values instead of formulas on row {row}.",
                        metadata.name,
                        row,
                    )
                records.append(RawRow(row, item.value, tuple(cell.value for cell in cells)))
            criteria.append(
                RawCriterion(metadata.name, tuple(header for _, header in raters), tuple(records))
            )
    finally:
        book.close()
    return tuple(criteria)
